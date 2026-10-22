"""数据驱动 —— 读取 CSV / Excel 参数化文件。

用例挂一个数据文件后按行执行：每行的「列名 → 值」作为运行时变量注入，
用例里用 ${列名} 引用（优先级高于全局变量）。

CSV 用 utf-8-sig 读（Excel「另存为 CSV」会带 BOM，不处理的话第一个列名会多出 \\ufeff）；
Excel 用 openpyxl 只读模式，首行当表头。
"""
import csv
import threading
from pathlib import Path

from openpyxl import load_workbook

from app.config import settings

CSV_SUFFIXES = {".csv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
ALLOWED_SUFFIXES = CSV_SUFFIXES | EXCEL_SUFFIXES


def dataset_dir(project_id: int) -> Path:
    path = Path(settings.dataset_dir) / str(project_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def resolve_path(project_id: int, filename: str) -> Path:
    """把文件名解析成数据目录下的路径，挡住 ../ 之类的越界访问与非法后缀。"""
    name = (filename or "").strip()
    if not name or Path(name).name != name or "/" in name or "\\" in name:
        raise ValueError("非法的文件名")
    if Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
        raise ValueError("只支持 .csv / .xlsx 文件")
    return dataset_dir(project_id) / name


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _load_csv(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            return []
        columns = [(c or "").strip() for c in reader.fieldnames]
        rows = []
        for raw in reader:
            # DictReader 的 key 是原始表头，这里按位置取值，避免表头带空格时对不上
            values = list(raw.values())
            rows.append({columns[i]: _clean(values[i]) for i in range(len(columns)) if columns[i]})
        return rows


def _load_excel(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            return []
        columns = [_clean(c) for c in header]
        rows = []
        for raw in iterator:
            values = list(raw)
            rows.append({columns[i]: _clean(values[i]) for i in range(len(columns)) if columns[i]})
        return rows
    finally:
        workbook.close()


def load_rows(project_id: int, filename: str) -> list[dict]:
    """读出一个数据文件的全部数据行，返回 [{列名: 值}]。

    全空行会被丢掉（Excel 里很容易留下尾部空行）。文件不存在抛 FileNotFoundError，
    格式非法抛 ValueError —— 都由调用方转成 HTTP 400。
    """
    path = resolve_path(project_id, filename)
    if not path.exists():
        raise FileNotFoundError(f"数据文件不存在：{filename}")

    try:
        rows = _load_csv(path) if path.suffix.lower() in CSV_SUFFIXES else _load_excel(path)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"解析失败：{type(e).__name__}: {str(e)[:120]}") from e

    return [r for r in rows if any(v for v in r.values())]


# 元信息缓存：(绝对路径, mtime_ns, size) → {rows, columns}
# 列表接口对每个文件调一次 describe()，而 describe() 原先要全量解析整个文件 ——
# 5 个 5000 行的 Excel 就是每刷一次列表解析 5 次。行数/列名不变就不该重解析。
#
# 用 (mtime, size) 当键而不是文件内容：同一秒内改两次且大小不变的情况极罕见，
# 真的撞上了也只是元信息晚一轮刷新，执行时读的仍是文件本身（load_rows 不缓存）。
_meta_cache: dict[tuple[str, int, int], dict] = {}
_meta_lock = threading.Lock()
_META_CACHE_MAX = 256


def invalidate(project_id: int, filename: str | None = None) -> None:
    """上传/删除后主动失效；filename 为空时清掉整个项目的缓存。"""
    root = Path(settings.dataset_dir) / str(project_id)
    target = root / filename if filename else None
    with _meta_lock:
        for key in list(_meta_cache):
            path = Path(key[0])
            if path.parent == root and (target is None or path == target):
                _meta_cache.pop(key, None)


def describe(project_id: int, filename: str) -> dict:
    """给列表/预览用的元信息（带 mtime+size 缓存，文件变了才重解析）。"""
    path = resolve_path(project_id, filename)
    stat = path.stat()  # 文件不存在时这里抛 FileNotFoundError，与原先 load_rows 的行为一致
    key = (str(path), stat.st_mtime_ns, stat.st_size)

    with _meta_lock:
        cached = _meta_cache.get(key)

    if cached is None:
        rows = load_rows(project_id, filename)
        columns: list[str] = []
        for r in rows:
            for column in r:
                if column not in columns:
                    columns.append(column)
        cached = {"rows": len(rows), "columns": columns}
        with _meta_lock:
            if len(_meta_cache) >= _META_CACHE_MAX:
                # 元信息本身很小，直接清空重建比维护 LRU 简单得多
                _meta_cache.clear()
            _meta_cache[key] = cached

    return {
        "filename": path.name,
        "size": stat.st_size,
        "rows": cached["rows"],
        "columns": cached["columns"],
    }
